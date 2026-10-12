# -*- coding: utf-8 -*-
"""
Hindi Order Maker v4 — Upload or paste/edit order format + Assessee Source Files

Install:
    py -m pip install -r requirements.txt
Run:
    py Hindi_Order_Maker.py

Privacy: files are processed locally on your PC for text/OCR extraction.
"""
import os, re, shutil, subprocess, json
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    from docx import Document
except ImportError:
    Document = None

APP_TITLE = "Hindi Order Maker v4"

def ocr_image(path):
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path), lang="hin+eng")
    except ImportError:
        raise RuntimeError("OCR packages missing. Run: py -m pip install Pillow pytesseract")
    except Exception as e:
        raise RuntimeError("Image OCR failed. Install Tesseract OCR and Hindi (hin) language data.\n\n" + str(e))

def ocr_pdf(path):
    try:
        import fitz
        import pytesseract
        from PIL import Image
        from io import BytesIO
    except ImportError:
        raise RuntimeError("Scanned PDF OCR packages missing. Run: py -m pip install pymupdf Pillow pytesseract")
    try:
        pdf = fitz.open(path)
        pages = []
        for i, page in enumerate(pdf, 1):
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            img = Image.open(BytesIO(pix.tobytes("png")))
            pages.append(f"--- Page {i} ---\n" + pytesseract.image_to_string(img, lang="hin+eng"))
        return "\n\n".join(pages)
    except Exception as e:
        raise RuntimeError("Scanned PDF OCR failed. Check Tesseract and Hindi language data.\n\n" + str(e))

def read_file(path):
    ext = Path(path).suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        return Path(path).read_text(encoding="utf-8-sig", errors="replace")
    if ext == ".pdf":
        text = ""
        if PdfReader:
            try:
                text = "\n\n".join(p.extract_text() or "" for p in PdfReader(path).pages).strip()
            except Exception:
                pass
        return text if text else ocr_pdf(path)
    if ext == ".docx":
        if not Document:
            raise RuntimeError("Install python-docx.")
        doc = Document(path)
        out = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                out.append(" | ".join(c.text for c in row.cells))
        return "\n".join(out)
    if ext in (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"):
        return ocr_image(path)
    raise ValueError("Supported: TXT, PDF, DOCX, CSV, PNG, JPG, TIFF")

def placeholders(text):
    return list(dict.fromkeys(re.findall(r"\{\{\s*([^{}]+?)\s*\}\}", text)))

def write_docx(text, path):
    if not Document:
        raise RuntimeError("Install python-docx: py -m pip install python-docx")
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(path)

def convert_pdf(docx_path, pdf_path):
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if lo:
        result = subprocess.run([lo, "--headless", "--convert-to", "pdf", "--outdir", str(Path(pdf_path).parent), docx_path],
                                capture_output=True, text=True, timeout=120)
        made = str(Path(pdf_path).parent / (Path(docx_path).stem + ".pdf"))
        if result.returncode == 0 and os.path.exists(made):
            if os.path.abspath(made) != os.path.abspath(pdf_path):
                if os.path.exists(pdf_path): os.remove(pdf_path)
                shutil.move(made, pdf_path)
            return
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        if os.path.exists(pdf_path): return
    except Exception:
        pass
    raise RuntimeError("PDF conversion needs Microsoft Word or LibreOffice. A DOCX copy has been saved beside the target.")

class App:
    def __init__(self, root):
        self.root = root
        root.title(APP_TITLE)
        root.geometry("1160x800")
        root.minsize(920, 650)
        self.template_path = ""
        self.case_files = []
        self.case_text = ""
        self.vars = {}
        style = ttk.Style()
        try: style.theme_use("clam")
        except Exception: pass
        style.configure("TLabel", font=("Segoe UI", 10))
        style.configure("TButton", font=("Segoe UI", 10))
        style.configure("TNotebook.Tab", font=("Segoe UI", 10), padding=(10, 6))
        style.configure("Heading.TLabel", font=("Segoe UI", 15, "bold"))
        head = ttk.Frame(root, padding=12); head.pack(fill="x")
        ttk.Label(head, text="Hindi Order Maker v4", style="Heading.TLabel").pack(side="left")
        ttk.Label(head, text="Upload or paste/edit a format + case details", foreground="#555").pack(side="right")
        self.nb = ttk.Notebook(root); self.nb.pack(fill="both", expand=True, padx=10, pady=(0, 8))
        self.t_template = ttk.Frame(self.nb, padding=10)
        self.t_case = ttk.Frame(self.nb, padding=10)
        self.t_fields = ttk.Frame(self.nb, padding=10)
        self.t_output = ttk.Frame(self.nb, padding=10)
        self.nb.add(self.t_template, text="1. Order format")
        self.nb.add(self.t_case, text="2. Assessee details")
        self.nb.add(self.t_fields, text="3. Review extracted fields")
        self.nb.add(self.t_output, text="4. Preview & export")
        self.build_template(); self.build_case(); self.build_fields(); self.build_output()
        ttk.Label(root, text="Review all extracted details and legal content before issuing any order.",
                  foreground="#8a3d00", padding=(10, 3)).pack(fill="x")

    def build_template(self):
        f=self.t_template
        bar=ttk.Frame(f); bar.pack(fill="x")
        ttk.Button(bar,text="Upload format (PDF/TXT/DOCX)",command=self.load_template).pack(side="left")
        ttk.Button(bar,text="Paste from clipboard",command=self.paste_template).pack(side="left",padx=6)
        ttk.Button(bar,text="Clear format",command=self.clear_template).pack(side="left")
        self.tpl_label=ttk.Label(bar,text="No format loaded — paste or type below"); self.tpl_label.pack(side="left",padx=10)
        ttk.Button(bar,text="Detect {{fields}}",command=self.refresh_fields).pack(side="right")
        ttk.Label(f,text="Option 1: Upload a PDF, TXT, or DOCX format. Option 2: copy a format and click 'Paste from clipboard', or paste/type directly into the editable box below. Make changes here before using the format. Scanned PDFs are OCR-converted to text; original page layout is not preserved. Use {{field_name}} placeholders for details that change for each assessee.", wraplength=1000).pack(anchor="w",pady=(10,5))
        self.tpl=tk.Text(f,wrap="word",font=("Segoe UI",11),undo=True); self.tpl.pack(fill="both",expand=True)
        self.tpl.insert("1.0","ORDER NO.: {{order_number}}\nDATE: {{date}}\n\nTo,\n{{assessee_name}}\n{{address}}\n\nSubject: {{subject}}\nAssessment Year: {{assessment_year}}\n\n[Insert order text here.]\n\n{{officer_name}}\n{{designation}}")

    def build_case(self):
        f=self.t_case
        bar=ttk.Frame(f); bar.pack(fill="x")
        ttk.Button(bar,text="Upload case details (multiple files)",command=self.load_case_files).pack(side="left")
        ttk.Button(bar,text="Extract text from selected files",command=self.extract_case).pack(side="left",padx=8)
        ttk.Button(bar,text="Clear files/text",command=self.clear_case).pack(side="left")
        self.case_label=ttk.Label(f,text="No case files selected"); self.case_label.pack(anchor="w",pady=8)
        ttk.Label(f,text="Extracted text from assessee-specific source documents. You can edit this text before field extraction.",
                  wraplength=1000).pack(anchor="w")
        self.case_box=tk.Text(f,wrap="word",font=("Segoe UI",10),undo=True); self.case_box.pack(fill="both",expand=True,pady=(5,0))

    def build_fields(self):
        f=self.t_fields
        bar=ttk.Frame(f); bar.pack(fill="x")
        ttk.Label(bar,text="Review the detected values and correct anything OCR missed.").pack(side="left")
        ttk.Button(bar,text="Extract likely details from case text",command=self.extract_fields).pack(side="right")
        self.fields_canvas=tk.Canvas(f,highlightthickness=0)
        scroll=ttk.Scrollbar(f,orient="vertical",command=self.fields_canvas.yview)
        self.fields_inner=ttk.Frame(self.fields_canvas)
        self.fields_inner.bind("<Configure>",lambda e:self.fields_canvas.configure(scrollregion=self.fields_canvas.bbox("all")))
        self.fields_canvas.create_window((0,0),window=self.fields_inner,anchor="nw")
        self.fields_canvas.configure(yscrollcommand=scroll.set)
        self.fields_canvas.pack(side="left",fill="both",expand=True,pady=10); scroll.pack(side="right",fill="y",pady=10)

    def build_output(self):
        f=self.t_output
        ttk.Label(f,text="Final order preview",style="Heading.TLabel").pack(anchor="w")
        self.preview=tk.Text(f,wrap="word",font=("Segoe UI",11)); self.preview.pack(fill="both",expand=True,pady=8)
        row=ttk.Frame(f); row.pack(fill="x")
        ttk.Button(row,text="Update preview",command=self.update_preview).pack(side="left")
        ttk.Button(row,text="Save Word (.docx)",command=self.save_docx).pack(side="left",padx=8)
        ttk.Button(row,text="Save PDF (.pdf)",command=self.save_pdf).pack(side="left")

    def paste_template(self):
        try:
            text = self.root.clipboard_get()
            if not text.strip():
                messagebox.showwarning("Clipboard empty", "Copy your order format first, then click Paste from clipboard.")
                return
            self.tpl.delete("1.0", "end")
            self.tpl.insert("1.0", text)
            self.template_path = ""
            self.tpl_label.config(text="Pasted format — editable")
            self.refresh_fields()
        except tk.TclError:
            messagebox.showwarning("Clipboard unavailable", "Copy your order format first, then click Paste from clipboard.")

    def clear_template(self):
        if messagebox.askyesno("Clear format", "Clear the current order format? This cannot be undone."):
            self.tpl.delete("1.0", "end")
            self.template_path = ""
            self.tpl_label.config(text="Blank format — paste or type below")
            self.refresh_fields()

    def load_template(self):
        path=filedialog.askopenfilename(title="Select existing format",filetypes=[("Supported","*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff *.txt *.md *.docx"),("All files","*.*")])
        if not path:return
        try:
            text=read_file(path)
            if not text.strip(): raise RuntimeError("No text was extracted from this file.")
            self.template_path=path; self.tpl.delete("1.0","end"); self.tpl.insert("1.0",text)
            self.tpl_label.config(text=Path(path).name + " — editable"); self.refresh_fields()
            messagebox.showinfo("Loaded","Format text loaded. Add {{placeholders}} to all parts that should change for each assessee.")
        except Exception as e: messagebox.showerror("Could not load format",str(e))

    def load_case_files(self):
        paths=filedialog.askopenfilenames(title="Select assessee case files",filetypes=[("Supported","*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff *.txt *.md *.docx *.csv"),("All files","*.*")])
        if paths:
            self.case_files=list(paths)
            self.case_label.config(text=f"{len(self.case_files)} file(s): "+", ".join(Path(p).name for p in self.case_files[:4]) + (" ..." if len(self.case_files)>4 else ""))

    def extract_case(self):
        if not self.case_files:
            messagebox.showwarning("No files","Upload one or more case-specific detail files first."); return
        parts=[]
        try:
            for p in self.case_files:
                parts.append(f"===== SOURCE: {Path(p).name} =====\n{read_file(p)}")
            self.case_text="\n\n".join(parts)
            self.case_box.delete("1.0","end"); self.case_box.insert("1.0",self.case_text)
            self.extract_fields()
            messagebox.showinfo("Extraction complete","Text extracted. Review the text and detected fields; OCR may make mistakes.")
        except Exception as e: messagebox.showerror("Extraction failed",str(e))

    def clear_case(self):
        self.case_files=[]; self.case_text=""; self.case_label.config(text="No case files selected")
        self.case_box.delete("1.0","end")

    def refresh_fields(self):
        names=placeholders(self.tpl.get("1.0","end-1c"))
        old={k:v.get() for k,v in self.vars.items()}
        for w in self.fields_inner.winfo_children():w.destroy()
        self.vars={}
        if not names:
            ttk.Label(self.fields_inner,text="No {{placeholders}} found. Add them to the format first.").grid(row=0,column=0,sticky="w")
            return
        ttk.Label(self.fields_inner,text="Placeholder",font=("Segoe UI",10,"bold")).grid(row=0,column=0,sticky="w",padx=5,pady=5)
        ttk.Label(self.fields_inner,text="Value",font=("Segoe UI",10,"bold")).grid(row=0,column=1,sticky="w",padx=5,pady=5)
        for i,n in enumerate(names,1):
            ttk.Label(self.fields_inner,text="{{"+n+"}}").grid(row=i,column=0,sticky="w",padx=5,pady=4)
            var=tk.StringVar(value=old.get(n,""))
            self.vars[n]=var
            ttk.Entry(self.fields_inner,textvariable=var,width=80).grid(row=i,column=1,sticky="ew",padx=5,pady=4)
        self.fields_inner.columnconfigure(1,weight=1)

    def extract_fields(self):
        text=self.case_box.get("1.0","end-1c")
        if not text.strip():
            messagebox.showwarning("No source text","Upload case files and click Extract text first."); return
        # Conservative, transparent heuristics: only autofill when a recognizable label is found.
        patterns={
            "assessee_name":[r"(?:assessee|taxpayer|नाम|करदाता)\s*[:：\-]\s*([^\n,;]+)"],
            "address":[r"(?:address|पता)\s*[:：\-]\s*([^\n]+)"],
            "assessment_year":[r"(?:assessment\s*year|निर्धारण\s*वर्ष|आकलन\s*वर्ष)\s*[:：\-]\s*([^\n,;]+)"],
            "order_number":[r"(?:order\s*(?:no|number)|आदेश\s*संख्या)\s*[:：\-]\s*([^\n,;]+)"],
            "date":[r"(?:date|दिनांक)\s*[:：\-]\s*([^\n,;]+)"],
            "PAN":[r"\b([A-Z]{5}[0-9]{4}[A-Z])\b"]
        }
        extracted={}
        for key, regs in patterns.items():
            for reg in regs:
                m=re.search(reg,text,re.I)
                if m:
                    extracted[key]=m.group(1).strip() if m.lastindex else m.group(0)
                    break
        # Map detected values to similarly named placeholders.
        for name,var in self.vars.items():
            normalized=re.sub(r"[^a-z0-9]", "", name.lower())
            aliases={
                "assessee_name":["assesseename","name","taxpayer","नाम","करदाता"],
                "address":["address","पता"],
                "assessment_year":["assessmentyear","year","निर्धारणवर्ष","आकलनवर्ष"],
                "order_number":["ordernumber","orderno","आदेशसंख्या"],
                "date":["date","दिनांक"],
                "pan":["pan"]
            }
            for k,val in extracted.items():
                if normalized in aliases.get(k,[]) or name.lower()==k.lower():
                    var.set(val)
                    break
        messagebox.showinfo("Field scan",f"Recognized {len(extracted)} possible labeled value(s): "+(", ".join(extracted.keys()) or "none")+
                            ". Values are only suggested by simple label matching; please verify them.")

    def rendered(self):
        text=self.tpl.get("1.0","end-1c")
        for name,var in self.vars.items():
            text=re.sub(r"\{\{\s*"+re.escape(name)+r"\s*\}\}",lambda m:var.get(),text)
        return text

    def update_preview(self):
        self.preview.delete("1.0","end"); self.preview.insert("1.0",self.rendered()); self.nb.select(self.t_output)

    def save_docx(self):
        path=filedialog.asksaveasfilename(defaultextension=".docx",initialfile="Order.docx",filetypes=[("Word document","*.docx")])
        if not path:return
        try:
            write_docx(self.rendered(),path); messagebox.showinfo("Saved","Word file saved:\n"+path)
        except Exception as e:messagebox.showerror("Save failed",str(e))

    def save_pdf(self):
        path=filedialog.asksaveasfilename(defaultextension=".pdf",initialfile="Order.pdf",filetypes=[("PDF document","*.pdf")])
        if not path:return
        docx=str(Path(path).with_suffix(".docx"))
        try:
            write_docx(self.rendered(),docx); convert_pdf(docx,path)
            messagebox.showinfo("Saved","PDF saved:\n"+path)
        except Exception as e:messagebox.showerror("PDF export failed",str(e))

def main():
    root=tk.Tk(); App(root); root.mainloop()

if __name__=="__main__":
    main()
